"""Comprehensive pytest suite for Resume-to-Job Description Matcher.

Covers all required test cases:
- TC1: Resume has Python, SQL, Pandas -> all three matched
- TC2: No required skills in resume -> 0%
- TC3: Every required skill present -> 100%
- TC4: Python mentioned five times -> counted once
- TC5: Empty resume -> validation error
- TC6: Alias (sklearn) -> normalized to scikit-learn
- TC7: Extra unrelated skills -> not counted as matches
- Main example -> exactly the expected output: score 75, matched: Python, SQL, Pandas, missing: Machine Learning
- Extra: Same resume as PDF, DOCX, TXT gives identical skills and score
- Extra: DOCX with skills in a table parses correctly
- Extra: Text-based PDF parses correctly
- Extra: Empty/corrupted PDF and unsupported file types give clear friendly errors
- Extra: Special symbols (C++, C#, .NET, Node.js) and Java vs JavaScript boundary handling
"""
import io
import pytest
import docx
from reportlab.pdfgen import canvas

from matcher.extract import SkillExtractor
from matcher.scoring import calculate_match
from matcher.parsers import (
    extract_text_from_file,
    validate_raw_text,
    parse_pdf,
    parse_docx,
    parse_txt,
    ValidationError,
    FileTooLargeError,
    UnsupportedFormatError,
    EmptyDocumentError,
    NoExtractableTextError,
    ParsingError,
)
from matcher.normalize import clean_text


@pytest.fixture(scope="module")
def extractor():
    """Shared SkillExtractor instance with skills.json."""
    return SkillExtractor()


# -------------------------------------------------------------------------
# Core Matching & Scoring Test Cases
# -------------------------------------------------------------------------

def test_main_example(extractor):
    """Main prompt example:
    JD: Python, SQL, Machine Learning, Pandas
    Resume: Python, Pandas, SQL, Excel
    Expected:
    {
      "matched_skills": ["Python", "SQL", "Pandas"],
      "missing_skills": ["Machine Learning"],
      "match_score": 75
    }
    """
    jd_text = "We are hiring a candidate with Python, SQL, Machine Learning, and Pandas."
    resume_text = "Skilled in Python, Pandas, SQL, and Excel for analytics."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    result = calculate_match(jd_skills, resume_skills)

    assert result["matched_skills"] == ["Python", "SQL", "Pandas"]
    assert result["missing_skills"] == ["Machine Learning"]
    assert result["match_score"] == 75
    assert "Excel" in result["additional_skills"]


def test_tc1_all_three_matched(extractor):
    """TC1: Resume has Python, SQL, Pandas -> all three matched."""
    jd_text = "Requirements: Python, SQL, and Pandas."
    resume_text = "Experience with Pandas, SQL, and Python programming."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    result = calculate_match(jd_skills, resume_skills)

    assert result["matched_skills"] == ["Python", "SQL", "Pandas"]
    assert result["missing_skills"] == []
    assert result["match_score"] == 100


def test_tc2_no_required_skills(extractor):
    """TC2: No required skills in resume -> 0%."""
    jd_text = "Job requires Python, SQL, and Machine Learning."
    resume_text = "Background in Excel, Tableau, and Graphic Design."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    result = calculate_match(jd_skills, resume_skills)

    assert result["matched_skills"] == []
    assert result["missing_skills"] == ["Python", "SQL", "Machine Learning"]
    assert result["match_score"] == 0


def test_tc3_every_required_skill_present(extractor):
    """TC3: Every required skill present -> 100%."""
    jd_text = "Must know Docker, Kubernetes, AWS."
    resume_text = "Expert in AWS cloud infrastructure, Kubernetes clusters, and Docker containers."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    result = calculate_match(jd_skills, resume_skills)

    assert result["matched_skills"] == ["Docker", "Kubernetes", "AWS"]
    assert result["missing_skills"] == []
    assert result["match_score"] == 100


def test_tc4_duplicate_mentions_counted_once(extractor):
    """TC4: Python mentioned five times -> counted once."""
    jd_text = "Required: Python."
    resume_text = (
        "I love Python. Python is great for scripting. I write Python daily. "
        "Python data structures are clean. Python projects."
    )

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    assert resume_skills == ["Python"]
    assert len(resume_skills) == 1

    result = calculate_match(jd_skills, resume_skills)
    assert result["matched_skills"] == ["Python"]
    assert result["match_score"] == 100
    assert result["matched_count"] == 1
    assert result["total_required"] == 1


def test_tc5_empty_resume_validation():
    """TC5: Empty resume -> validation error."""
    with pytest.raises(EmptyDocumentError) as exc_info:
        validate_raw_text("", document_name="Resume")
    assert "Resume is empty" in str(exc_info.value)

    with pytest.raises(EmptyDocumentError) as exc_info2:
        validate_raw_text("   \n\t   ", document_name="Resume")
    assert "Resume is empty or contains only whitespace" in str(exc_info2.value)


def test_tc6_alias_normalization(extractor):
    """TC6: Alias (sklearn) -> normalized to scikit-learn."""
    jd_text = "Job description specifies scikit-learn and Machine Learning."
    resume_text = "Worked with sklearn and ML algorithms in production."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    assert "scikit-learn" in resume_skills
    assert "Machine Learning" in resume_skills

    result = calculate_match(jd_skills, resume_skills)
    assert result["matched_skills"] == ["scikit-learn", "Machine Learning"]
    assert result["missing_skills"] == []
    assert result["match_score"] == 100


def test_tc7_extra_unrelated_skills_not_counted(extractor):
    """TC7: Extra unrelated skills -> not counted as matches, separated in additional_skills."""
    jd_text = "Requirements: Python, SQL."
    resume_text = "Proficient in Python, SQL, MongoDB, Redis, Docker, and Tableau."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    result = calculate_match(jd_skills, resume_skills)

    assert result["matched_skills"] == ["Python", "SQL"]
    assert result["match_score"] == 100
    assert set(result["additional_skills"]) == {"MongoDB", "Redis", "Docker", "Tableau"}
    # Verify results DataFrame shows correct status
    df = result["results_df"]
    mongo_row = df[df["Skill"] == "MongoDB"].iloc[0]
    assert mongo_row["Status"] == "Additional (Not in JD)"
    assert mongo_row["In JD"] == "No"
    assert mongo_row["In Resume"] == "Yes"


# -------------------------------------------------------------------------
# Special Programming Term Boundaries & Distinctions
# -------------------------------------------------------------------------

def test_java_vs_javascript(extractor):
    """Java must not match JavaScript."""
    jd_text = "Looking for a Core Java developer."
    resume_text_js_only = "Senior JavaScript frontend developer with React."
    resume_text_both = "Full-stack developer with Java and JavaScript."

    jd_skills = extractor.extract_skills(jd_text)
    assert jd_skills == ["Java"]

    # JS only candidate
    res_skills_js = extractor.extract_skills(resume_text_js_only)
    assert "JavaScript" in res_skills_js
    assert "Java" not in res_skills_js
    res_js_match = calculate_match(jd_skills, res_skills_js)
    assert res_js_match["matched_skills"] == []
    assert res_js_match["match_score"] == 0

    # Both candidate
    res_skills_both = extractor.extract_skills(resume_text_both)
    assert "Java" in res_skills_both
    assert "JavaScript" in res_skills_both
    res_both_match = calculate_match(jd_skills, res_skills_both)
    assert res_both_match["matched_skills"] == ["Java"]
    assert res_both_match["match_score"] == 100


def test_special_tokens_cpp_csharp_dotnet_nodejs(extractor):
    """Ensure C++, C#, .NET, Node.js parse and match cleanly without colliding with C or Node."""
    jd_text = "Need C++, C#, .NET, and Node.js engineers."
    resume_text = "Experienced in C++, Node.js, C#, and .NET development."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    assert jd_skills == ["C++", "C#", ".NET", "Node.js"]
    assert set(resume_skills) == {"C++", "C#", ".NET", "Node.js"}

    result = calculate_match(jd_skills, resume_skills)
    assert result["match_score"] == 100
    assert result["matched_skills"] == ["C++", "C#", ".NET", "Node.js"]


def test_zero_recognized_skills_in_jd(extractor):
    """JD with zero recognized skills -> handle gracefully, no division by zero."""
    jd_text = "Looking for an energetic, collaborative, and hard-working team player."
    resume_text = "Python, SQL, Machine Learning expert."

    jd_skills = extractor.extract_skills(jd_text)
    resume_skills = extractor.extract_skills(resume_text)

    assert jd_skills == []
    result = calculate_match(jd_skills, resume_skills)
    assert result["match_score"] == 0
    assert result["total_required"] == 0
    assert result["matched_count"] == 0
    assert "No recognized technical skills detected" in result["recommendation"]


def test_clean_text_hyphenated_line_breaks():
    """Verify hyphenated line breaks are repaired seamlessly."""
    raw = "Hands-on experience in ma-\nchine learning and sci-\nkit-learn."
    cleaned = clean_text(raw)
    assert "machine learning" in cleaned
    assert "scikit learn" in cleaned or "scikitlearn" in cleaned or "scikit" in cleaned


# -------------------------------------------------------------------------
# Document Parsers & Cross-Format Equivalence Tests
# -------------------------------------------------------------------------

def _create_test_pdf(text: str) -> bytes:
    bio = io.BytesIO()
    c = canvas.Canvas(bio)
    c.drawString(72, 750, text)
    c.save()
    bio.seek(0)
    return bio.getvalue()


def _create_test_docx(paragraphs: list, table_data: list = None) -> bytes:
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = str(val)
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()


def test_identical_skills_across_pdf_docx_txt(extractor):
    """Extra: Same resume content in PDF, DOCX, and TXT gives identical matched skills and score."""
    sample_content = "Proficient in Python, Pandas, SQL, and Excel for statistical analysis."
    jd_content = "Looking for a candidate with Python, SQL, Machine Learning, and Pandas."

    pdf_bytes = _create_test_pdf(sample_content)
    docx_bytes = _create_test_docx([sample_content])
    txt_bytes = sample_content.encode("utf-8")

    text_pdf = extract_text_from_file(pdf_bytes, filename="resume.pdf")
    text_docx = extract_text_from_file(docx_bytes, filename="resume.docx")
    text_txt = extract_text_from_file(txt_bytes, filename="resume.txt")

    skills_pdf = extractor.extract_skills(text_pdf)
    skills_docx = extractor.extract_skills(text_docx)
    skills_txt = extractor.extract_skills(text_txt)

    assert skills_pdf == skills_docx == skills_txt == ["Python", "Pandas", "SQL", "Excel"]

    jd_skills = extractor.extract_skills(jd_content)

    res_pdf = calculate_match(jd_skills, skills_pdf)
    res_docx = calculate_match(jd_skills, skills_docx)
    res_txt = calculate_match(jd_skills, skills_txt)

    assert res_pdf["match_score"] == res_docx["match_score"] == res_txt["match_score"] == 75
    assert res_pdf["matched_skills"] == res_docx["matched_skills"] == res_txt["matched_skills"] == ["Python", "SQL", "Pandas"]
    assert res_pdf["missing_skills"] == res_docx["missing_skills"] == res_txt["missing_skills"] == ["Machine Learning"]


def test_docx_table_parsing(extractor):
    """Extra: DOCX with skills formatted in a table parses correctly."""
    docx_bytes = _create_test_docx(
        paragraphs=["Summary: Data Engineer"],
        table_data=[
            ["Programming", "Python, SQL"],
            ["Frameworks", "Pandas, Fastapi"],
            ["Cloud", "AWS, Docker"]
        ]
    )

    extracted_text = extract_text_from_file(docx_bytes, filename="profile.docx")
    skills = extractor.extract_skills(extracted_text)

    assert "Python" in skills
    assert "SQL" in skills
    assert "Pandas" in skills
    assert "AWS" in skills
    assert "Docker" in skills


def test_unsupported_file_type():
    """Unsupported type (e.g. .jpg, .doc) -> friendly message listing supported formats."""
    with pytest.raises(UnsupportedFormatError) as exc_info:
        extract_text_from_file(b"image data", filename="resume.jpg")
    assert "Unsupported file format '.jpg'" in str(exc_info.value)
    assert "Supported formats are: PDF (.pdf), Word (.docx), and Plain Text (.txt)" in str(exc_info.value)


def test_file_too_large():
    """File exceeding 5 MB limit -> FileTooLargeError."""
    large_data = b"a" * (5 * 1024 * 1024 + 1024)
    with pytest.raises(FileTooLargeError) as exc_info:
        extract_text_from_file(large_data, filename="large.txt")
    assert "exceeds the maximum allowed limit of 5 MB" in str(exc_info.value)


def test_empty_or_corrupted_pdf():
    """Empty or corrupted PDF raises clear friendly error suggesting pasting text."""
    # Blank PDF with no text
    bio = io.BytesIO()
    c = canvas.Canvas(bio)
    c.showPage()
    c.save()
    blank_pdf_bytes = bio.getvalue()

    with pytest.raises(NoExtractableTextError) as exc_info:
        extract_text_from_file(blank_pdf_bytes, filename="scanned.pdf")
    assert "No extractable text found in this PDF" in str(exc_info.value)
    assert "Please copy and paste the text instead" in str(exc_info.value)

    # Corrupted bytes
    with pytest.raises(ParsingError) as exc_info2:
        extract_text_from_file(b"This is not a valid pdf header", filename="corrupted.pdf")
    assert "Could not open PDF file" in str(exc_info2.value)
