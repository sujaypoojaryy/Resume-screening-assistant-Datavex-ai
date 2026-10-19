"""File parsing and text extraction module.

Supports PDF (.pdf), Word (.docx), and plain text (.txt) files, as well as pasted text.
Handles validation (size, format, empty files, scanned PDFs) with clear user-friendly errors.
"""
from io import BytesIO
from pathlib import Path
from typing import Union, BinaryIO
import pypdf
import docx

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ParsingError(Exception):
    """Base exception for parsing and validation errors."""
    pass


class ValidationError(ParsingError):
    """Raised when input validation fails."""
    pass


class FileTooLargeError(ValidationError):
    """Raised when uploaded file exceeds size limit."""
    pass


class UnsupportedFormatError(ValidationError):
    """Raised when file extension is not supported."""
    pass


class EmptyDocumentError(ValidationError):
    """Raised when document contains no text."""
    pass


class NoExtractableTextError(ParsingError):
    """Raised when PDF has no extractable text (scanned, image-only, or encrypted)."""
    pass


def validate_raw_text(text: str, document_name: str = "Document") -> str:
    """Validate pasted text input.
    
    Args:
        text: The string to validate.
        document_name: Human-friendly name for error messages.
        
    Returns:
        The validated string if non-empty.
        
    Raises:
        EmptyDocumentError: If text is empty or contains only whitespace.
    """
    if text is None or not str(text).strip():
        raise EmptyDocumentError(
            f"{document_name} is empty or contains only whitespace. Please provide valid text content."
        )
    return str(text).strip()


def _get_bytes_and_name(source: Union[str, Path, BinaryIO, bytes], filename: str = None):
    """Normalize file source into (bytes_data, filename)."""
    if isinstance(source, (str, Path)):
        p = Path(source)
        if not p.exists():
            raise ValidationError(f"File not found: {p}")
        name = filename or p.name
        with open(p, "rb") as f:
            data = f.read()
        return data, name

    if isinstance(source, bytes):
        name = filename or "uploaded_file"
        return source, name

    # File-like object (e.g., io.BytesIO or Streamlit UploadedFile)
    name = filename or getattr(source, "name", "uploaded_file")
    if hasattr(source, "seek"):
        source.seek(0)
    data = source.read()
    if hasattr(source, "seek"):
        source.seek(0)
    return data, name


def parse_pdf(stream_or_bytes: Union[BinaryIO, bytes]) -> str:
    """Extract text page-by-page from a PDF document using pypdf.
    
    Handles encrypted, corrupted, and image-only/scanned PDFs gracefully.
    """
    bio = BytesIO(stream_or_bytes) if isinstance(stream_or_bytes, bytes) else stream_or_bytes
    try:
        reader = pypdf.PdfReader(bio)
    except Exception as e:
        raise ParsingError(
            f"Could not open PDF file. It may be corrupted or in an invalid format. "
            f"Please try pasting the text directly instead."
        ) from e

    if reader.is_encrypted:
        try:
            decrypted = reader.decrypt("")
            if decrypted == 0:
                raise NoExtractableTextError(
                    "This PDF is password-protected. Please remove password protection or paste the text directly."
                )
        except Exception:
            raise NoExtractableTextError(
                "This PDF is password-protected. Please remove password protection or paste the text directly."
            )

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text.strip())
        except Exception as e:
            # Continue reading other pages if one page extraction fails
            continue

    combined_text = "\n".join(extracted_pages).strip()
    if not combined_text:
        raise NoExtractableTextError(
            "No extractable text found in this PDF. It may be a scanned or image-only document. "
            "Please copy and paste the text instead."
        )

    return combined_text


def parse_docx(stream_or_bytes: Union[BinaryIO, bytes]) -> str:
    """Extract text from Word (.docx) documents, including paragraphs and tables."""
    bio = BytesIO(stream_or_bytes) if isinstance(stream_or_bytes, bytes) else stream_or_bytes
    try:
        doc = docx.Document(bio)
    except Exception as e:
        raise ParsingError(
            f"Could not read Word document (.docx). The file may be damaged or invalid. "
            f"Please try pasting the text directly instead."
        ) from e

    text_parts = []
    # Extract from paragraphs
    for paragraph in doc.paragraphs:
        txt = paragraph.text.strip()
        if txt:
            text_parts.append(txt)

    # Extract from tables (rows and cells)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    combined = "\n".join(text_parts).strip()
    if not combined:
        raise EmptyDocumentError(
            "The Word document (.docx) does not contain any readable text. "
            "Please check the file or paste the text directly."
        )

    return combined


def parse_txt(data: bytes) -> str:
    """Extract text from plain text files, trying UTF-8 with fallback to Latin-1."""
    if not data:
        raise EmptyDocumentError("The plain text file is empty. Please provide a file with content.")
    
    try:
        content = data.decode("utf-8")
    except UnicodeDecodeError:
        content = data.decode("latin-1", errors="replace")

    trimmed = content.strip()
    if not trimmed:
        raise EmptyDocumentError("The plain text file contains only whitespace. Please provide valid text content.")

    return trimmed


def extract_text_from_file(source: Union[str, Path, BinaryIO, bytes], filename: str = None) -> str:
    """Route document to the proper parser based on file extension.
    
    Validates:
      - Max file size (5 MB limit)
      - Supported extensions (.pdf, .docx, .txt)
      - Empty or unreadable files
      
    Returns:
        Extracted raw text string.
    """
    raw_bytes, name = _get_bytes_and_name(source, filename)

    # Check file size
    if len(raw_bytes) > MAX_FILE_SIZE_BYTES:
        mb_size = round(len(raw_bytes) / (1024 * 1024), 2)
        raise FileTooLargeError(
            f"File '{name}' ({mb_size} MB) exceeds the maximum allowed limit of 5 MB. "
            f"Please upload a smaller file."
        )

    # Check extension
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported_list = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFormatError(
            f"Unsupported file format '{suffix}' for '{name}'. "
            f"Supported formats are: PDF (.pdf), Word (.docx), and Plain Text (.txt)."
        )

    # Route parser
    if suffix == ".pdf":
        return parse_pdf(raw_bytes)
    elif suffix == ".docx":
        return parse_docx(raw_bytes)
    elif suffix == ".txt":
        return parse_txt(raw_bytes)
    else:
        raise UnsupportedFormatError(f"No parser available for format: {suffix}")
